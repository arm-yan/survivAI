"""Evaluate the generated workbooks' formulas and check them against hand-calculated cases.

Usage (from the repo root, after build.py):
    .venv\\Scripts\\python.exe products\\seller-profit-tracker\\tests\\run_tests.py

How it works: for each "pass" the script copies the built workbook, replaces the sample
data with the test inputs from cases.md (openpyxl), evaluates every formula with the
`formulas` library (an independent Excel formula engine), and compares the results with
the hand-calculated expected values. Money must match within $0.01, percentages within
0.01 percentage points. It also runs static checks on every formula (allowed functions
only, no hardcoded rates).
"""

from __future__ import annotations

import datetime as dt
import logging
import os
import re
import shutil
import sys
import tempfile
import zipfile
from pathlib import Path

import formulas

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))
from build import FEE_ROW_OF  # noqa: E402
from openpyxl import load_workbook

logging.disable(logging.WARNING)
os.environ.setdefault("TQDM_DISABLE", "1")

HERE = Path(__file__).resolve().parent
DIST = HERE.parent / "dist"
TRACKER = DIST / "MarginKite-Seller-Profit-Tracker.xlsx"
LITE = DIST / "MarginKite-Etsy-Fee-Calculator-Lite.xlsx"
TRACKER_BLANK = DIST / "MarginKite-Seller-Profit-Tracker-BLANK.xlsx"
LITE_BLANK = DIST / "MarginKite-Etsy-Fee-Calculator-Lite-BLANK.xlsx"
ALL_FILES = (TRACKER, LITE, TRACKER_BLANK, LITE_BLANK)

MONEY_TOL = 0.01
PCT_TOL = 0.0001

results: list[tuple[str, str, bool, str]] = []


def check(case, what, actual, expected, tol=MONEY_TOL):
    if isinstance(expected, str):
        ok = actual == expected
    else:
        ok = isinstance(actual, (int, float)) and abs(actual - expected) <= tol
    results.append((case, what, ok, f"expected {expected!r}, got {actual!r}"))


# ------------------------------------------------------------------ evaluation
def evaluate(path: Path):
    xl = formulas.ExcelModel().loads(str(path)).finish()
    sol = xl.calculate()
    prefix = f"'[{path.name}]"
    out = {}
    for k, v in sol.items():
        if k.startswith(prefix) and ":" not in k:
            sheet, cell = k[len(prefix):].split("'!")
            val = v.value[0, 0] if hasattr(v, "value") else v
            if hasattr(val, "item"):
                val = val.item()
            out[(sheet, cell)] = val
    return lambda sheet, cell: _norm(out.get((sheet.upper(), cell)))


def _norm(v):
    if v is None:
        return None
    s = str(v)
    if s in ("empty", ""):
        return ""
    if isinstance(v, bool):
        return v
    try:
        return float(v)
    except (TypeError, ValueError):
        return s


# ------------------------------------------------------------------ test setup
TEST_PRODUCTS = [("Test Mug", 4.00, 1.00, 30), ("Test Print", 2.00, 0.50, 0)]

D = dt.date(2026, 9, 5)
# date, platform, product, qty, price, ship charged, ship cost, discount, offsite, refunded, partial, credited
ORDERS_A = {
    "C01": [D, "Etsy", "Test Mug", 1, 30, 5, 4, None, "N", "N", None, None],
    "C02": [D, "Etsy", "Test Mug", 1, 30, 0, 4, None, "N", "N", None, None],
    "C03": [D, "Etsy", "Test Mug", 1, 30, 5, 4, None, "Y", "N", None, None],
    "C04": [D, "Etsy", "Test Mug", 1, 800, 0, 10, None, "Y", "N", None, None],
    "C05": [D, "Etsy", "Test Mug", 1, 30, 5, 4, None, "N", "Y", None, "Y"],
    "C06": [D, "Etsy", "Test Mug", 1, 30, 5, 4, None, "N", "Y", 10, "Y"],
    "C07": [D, "Etsy", "Test Mug", 1, 30, 5, 4, None, "N", "Y", 10, "N"],
    "C08": [D, "Etsy", "Test Mug", 2, 30, 5, 6, 6, "N", "N", None, None],
    "C09": [D, "Shopify", "Test Mug", 1, 30, 5, 4, None, "N", "N", None, None],
}
ORDERS_C = {
    "C15": [D, "Etsy", "Test Mug", 1, 30, 5, 4, None, "Y", "N", None, None],
    "C16": [D, "Shopify", "Test Mug", 1, 30, 5, 4, None, "N", "N", None, None],
    "C19": [D, "Etsy", "Test Mug", 2, 30, 5, 6, 6, "N", "N", None, None],
}
TAX_CELL = f"C{FEE_ROW_OF['RULE_TAX_BASE']}"
DISCOUNT_RULE_CELL = f"C{FEE_ROW_OF['RULE_DISCOUNT_BASE']}"


def prepare(src: Path, dst: Path, settings: dict, orders: dict):
    wb = load_workbook(src)
    s = wb["Settings"]
    for cell, val in settings.items():
        s[cell] = val
    p = wb["Products"]
    for r in range(5, 55):
        for c in "ABCD":
            p[f"{c}{r}"] = None
    for i, row in enumerate(TEST_PRODUCTS):
        for c, val in zip("ABCD", row):
            p[f"{c}{5 + i}"] = val
    o = wb["Orders Log"]
    for r in range(6, 506):
        for c in range(1, 13):
            o.cell(row=r, column=c).value = None
    rows = {}
    for i, (case, data) in enumerate(orders.items()):
        r = 6 + i
        rows[case] = r
        for c, val in enumerate(data, 1):
            o.cell(row=r, column=c).value = val
    e = wb["Expenses"]
    for r in range(6, 66):
        for c in "ABCDEF":
            e[f"{c}{r}"] = None
    for c, val in zip("ABCDEF", [dt.date(2026, 9, 10), 10, 27, 5, 8, 0]):
        e[f"{c}6"] = val
    calc = wb["Price Calculator"]
    for cell, val in {"C4": "Test Mug", "C6": 30, "C7": 5, "C8": 4, "C9": "N", "C10": 50, "C11": 0.30}.items():
        calc[cell] = val
    wb["Dashboard"]["C4"] = dt.date(2026, 9, 1)
    wb.save(dst)
    return rows


BASE = {"C5": 20, "C6": "N", "C7": "Basic", "C8": "Monthly", "C9": "Shopify Payments"}


def order_checks(g, case, r, exp):
    names = {"M": "order total", "N": "listing fee", "O": "transaction fee", "P": "processing",
             "Q": "offsite ads fee", "R": "fee credit", "S": "total fees", "T": "refund", "U": "product cost",
             "V": "net profit", "W": "margin"}
    for col, val in exp.items():
        check(case, names[col], g("Orders Log", f"{col}{r}"), val, PCT_TOL if col == "W" else MONEY_TOL)


def run_pass_a(tmp: Path):
    dst = tmp / TRACKER.name
    rows = prepare(TRACKER, dst, BASE, ORDERS_A)
    g = evaluate(dst)
    check("C00 product unit cost", "Products labour cost (30 min x $20)", g("Products", "E5"), 10.00)
    check("C00 product unit cost", "Products unit cost", g("Products", "F5"), 15.00)
    order_checks(g, "C01 basic Etsy sale", rows["C01"],
                 {"M": 35, "N": 0.20, "O": 2.275, "P": 1.30, "Q": 0, "S": 3.775, "U": 15, "V": 12.225, "W": 0.349286})
    order_checks(g, "C02 free shipping", rows["C02"], {"M": 30, "O": 1.95, "P": 1.15, "S": 3.30, "V": 7.70})
    order_checks(g, "C03 Offsite Ads 15%", rows["C03"], {"Q": 5.25, "S": 9.025, "V": 6.975})
    order_checks(g, "C04 Offsite Ads $100 cap", rows["C04"],
                 {"M": 800, "O": 52, "P": 24.25, "Q": 100, "S": 176.45, "V": 598.55})
    order_checks(g, "C05 full refund", rows["C05"], {"T": 35, "R": 3.325, "S": 0.45, "V": -19.45})
    order_checks(g, "C06 partial refund, fees credited", rows["C06"], {"T": 10, "R": 0.95, "S": 2.825, "V": 3.175})
    order_checks(g, "C07 partial refund, no fee credit", rows["C07"], {"T": 10, "R": 0, "S": 3.775, "V": 2.225})
    order_checks(g, "C08 discount, qty 2", rows["C08"],
                 {"M": 59, "N": 0.40, "O": 3.835, "P": 2.02, "S": 6.255, "U": 30, "V": 16.745, "W": 0.283814})
    order_checks(g, "C09 Shopify order (Shopify Payments)", rows["C09"], {"P": 1.315, "S": 1.315, "V": 14.685})
    c = "Price Calculator"
    check("C10 Shopify sale, monthly billing", "Etsy profit", g(c, "C25"), 12.225)
    check("C10 Shopify sale, monthly billing", "plan cost per order", g(c, "D21"), 0.54)
    check("C10 Shopify sale, monthly billing", "Shopify total fees", g(c, "D22"), 1.855)
    check("C10 Shopify sale, monthly billing", "Shopify profit", g(c, "D25"), 14.145)
    check("C10 Shopify sale, monthly billing", "Shopify margin", g(c, "D26"), 0.404143, PCT_TOL)
    check("C12 reverse margin 30%", "Etsy price needed", g(c, "C31"), 27.15)
    check("C12 reverse margin 30%", "Shopify price needed (monthly)", g(c, "D31"), 24.57)
    check("C13 Shopify break-even", "saving per order", g(c, "C36"), 2.46)
    check("C13 Shopify break-even", "orders/month (monthly $27)", g(c, "C38"), 11, 0)
    d = "Dashboard"
    check("C17 dashboard roll-up", "Sep revenue", g(d, "D7"), 1099)
    check("C17 dashboard roll-up", "Sep orders", g(d, "C7"), 9, 0)
    check("C17 dashboard roll-up", "Sep overhead", g(d, "I7"), 50)
    check("C17 dashboard roll-up", "Sep net profit", g(d, "J7"), 592.83)
    check("C17 dashboard roll-up", "fee breakdown total = fees column", g(d, "D37"), g(d, "F19"))
    check("C17 dashboard roll-up", "top product", g(d, "B24"), "Test Mug")
    check("C17 dashboard roll-up", "top product profit", g(d, "D24"), 642.83)
    check("C17 dashboard roll-up", "bottom product", g(d, "H24"), "Test Print")


def run_pass_b(tmp: Path):
    dst = tmp / TRACKER.name
    rows = prepare(TRACKER, dst, {**BASE, "C8": "Yearly", TAX_CELL: 0.08}, ORDERS_A)
    g = evaluate(dst)
    case = "C18 sales tax 8% in processing base"
    order_checks(g, case, rows["C01"], {"O": 2.275, "P": 1.384, "S": 3.859, "V": 12.141})
    check(case, "calculator Etsy processing", g("Price Calculator", "C19"), 1.384)
    check(case, "calculator Etsy transaction (no tax)", g("Price Calculator", "C18"), 2.275)
    check(case, "Shopify processing unaffected", g("Price Calculator", "D19"), 1.315)
    c = "Price Calculator"
    case = "C11 Shopify sale, yearly billing"
    check(case, "plan cost per month", g("Settings", "C13"), 19)
    check(case, "plan cost per order", g(c, "D21"), 0.38)
    check(case, "Shopify total fees", g(c, "D22"), 1.695)
    check(case, "Shopify profit", g(c, "D25"), 14.305)
    check("C12 reverse margin 30%", "Shopify price needed (yearly)", g(c, "D31"), 24.33)
    check("C13 Shopify break-even", "orders/month (yearly $19)", g(c, "C38"), 8, 0)


def run_pass_c(tmp: Path):
    dst = tmp / TRACKER.name
    rows = prepare(TRACKER, dst, {**BASE, "C6": "Y", "C9": "Third-party provider", DISCOUNT_RULE_CELL: "N"}, ORDERS_C)
    g = evaluate(dst)
    check("C15 Offsite Ads 12% ($10k+ shop)", "rate in use", g("Settings", "C12"), 0.12, PCT_TOL)
    order_checks(g, "C15 Offsite Ads 12% ($10k+ shop)", rows["C15"], {"Q": 4.20, "S": 7.975, "V": 8.025})
    order_checks(g, "C16 Shopify, third-party provider", rows["C16"], {"P": 2.015, "V": 13.985})
    check("C16 Shopify, third-party provider", "calculator Shopify profit", g("Price Calculator", "D25"), 13.445)
    order_checks(g, "C19 discount, alternative rule N (pre-discount base)", rows["C19"],
                 {"M": 59, "O": 4.225, "P": 2.02, "S": 6.645, "V": 16.355})


def run_lite(tmp: Path):
    dst = tmp / LITE.name
    wb = load_workbook(LITE)
    ws = wb.active
    for cell, val in {"C5": 30, "C6": 5, "C7": 4, "C8": 15, "C9": "Y", "C10": "N"}.items():
        ws[cell] = val
    wb.save(dst)
    g = evaluate(dst)
    s = "Etsy Fee Calculator"
    case = "C14 Lite calculator (C01 + Offsite Ad)"
    for cell, what, val in [("C13", "order total", 35), ("C14", "listing", 0.20), ("C15", "transaction", 2.275),
                            ("C16", "processing", 1.30), ("C17", "offsite ads", 5.25), ("C18", "total fees", 9.025),
                            ("C20", "net profit", 6.975)]:
        check(case, what, g(s, cell), val)
    check(case, "margin", g(s, "C21"), 0.199286, PCT_TOL)


# ------------------------------------------------------------------ static checks
ALLOWED = {"IF", "AND", "OR", "MIN", "MAX", "SUM", "SUMIFS", "COUNTIFS", "INDEX", "MATCH", "IFERROR",
           "LARGE", "SMALL", "ROUNDUP", "DATE", "YEAR", "MONTH", "ROW", "TEXT", "ISNUMBER"}
ALLOWED_LITERALS = {"0", "1", "2", "3", "60", "1000000000"}
ALLOWED_URLS = {
    "https://docs.google.com/spreadsheets/d/1EC3WQBv85J-L67KSHyWd0U1QExs7YLxZFE5qH7lfojw/copy",
    "https://docs.google.com/spreadsheets/d/1I5_ppRbvAYHL0KJ4ljvBSgQvCpiCP-S8ohdihBVczeo/copy",
    "https://docs.google.com/spreadsheets/d/1f_ZIgOi_mZsbkiiI2cuk_XiTdiTpllCazWIeVIYob00/copy",
    "https://marginkite.gumroad.com/l/profit-tracker",
    "https://marginkite.gumroad.com/l/free-fee-calculator",
    "mailto:marginkite.shop@gmail.com",
}


def static_checks():
    for path in ALL_FILES:
        wb = load_workbook(path)
        n = 0
        bad_fn, bad_lit, ext = set(), set(), 0
        for ws in wb.worksheets:
            if not ws.protection.sheet or ws.protection.password:
                results.append(("S0", f"{path.name}/{ws.title} protected without password", False, "not protected or has password"))
            for row in ws.iter_rows():
                for cell in row:
                    v = cell.value
                    if isinstance(v, str) and v.startswith("="):
                        n += 1
                        body = re.sub(r'"[^"]*"', '""', v)
                        bad_fn |= set(re.findall(r"([A-Z][A-Z0-9.]*)\(", body)) - ALLOWED
                        refs_removed = re.sub(r"\$?[A-Z]{1,3}\$?\d+", "", body)
                        bad_lit |= set(re.findall(r"(?<![A-Za-z_])\d+(?:\.\d+)?", refs_removed)) - ALLOWED_LITERALS
                        if "[" in v or "http" in v.lower():
                            ext += 1
                        if not cell.protection.locked:
                            results.append(("S0", f"{ws.title}!{cell.coordinate} formula locked", False, "unlocked"))
        results.append(("S1", f"{path.name}: {n} formulas use only Excel 2016/Sheets functions", not bad_fn,
                        f"disallowed: {sorted(bad_fn)}"))
        results.append(("S2", f"{path.name}: no hardcoded rates in formulas", not bad_lit,
                        f"literals: {sorted(bad_lit)}"))
        # Only the approved MarginKite URLs may appear as hyperlinks; no external-workbook links at all.
        links = set()
        with zipfile.ZipFile(path) as z:
            ext_parts = [n for n in z.namelist() if "externalLink" in n]
            for name in z.namelist():
                if name.endswith(".rels"):
                    xml = z.read(name).decode("utf-8")
                    links |= set(re.findall(r'Target="([^"]+)"\s+TargetMode="External"', xml))
                    links |= set(re.findall(r'TargetMode="External"\s+Target="([^"]+)"', xml))
        bad_links = links - ALLOWED_URLS
        results.append(("S3", f"{path.name}: no external links except approved MarginKite URLs {sorted(links)}",
                        ext == 0 and not wb._external_links and not ext_parts and not bad_links,
                        f"{ext} formulas with URLs/workbook refs; external parts {ext_parts}; unapproved {sorted(bad_links)}"))


def shipped_no_errors():
    """The files as shipped (with sample data) must evaluate with no #errors in any formula cell."""
    for path in ALL_FILES:
        g = evaluate(path)
        wb = load_workbook(path)
        bad = []
        for ws in wb.worksheets:
            for row in ws.iter_rows():
                for cell in row:
                    if isinstance(cell.value, str) and cell.value.startswith("="):
                        v = g(ws.title, cell.coordinate)
                        if isinstance(v, str) and v.startswith("#"):
                            bad.append(f"{ws.title}!{cell.coordinate}={v}")
        case = "S5" if "BLANK" in path.name else "S4"
        results.append((case, f"{path.name}: every formula cell evaluates with no errors", not bad, f"{bad[:5]}"))


def main():
    if not all(p.exists() for p in ALL_FILES):
        sys.exit("Run build.py first.")
    static_checks()
    shipped_no_errors()
    with tempfile.TemporaryDirectory() as t:
        tmp = Path(t)
        for fn in (run_pass_a, run_pass_b, run_pass_c, run_lite):
            fn(tmp)
            for f in tmp.iterdir():
                f.unlink() if f.is_file() else shutil.rmtree(f)

    by_case: dict[str, list] = {}
    for case, what, ok, msg in results:
        by_case.setdefault(case, []).append((what, ok, msg))
    failed = 0
    for case in sorted(by_case):
        items = by_case[case]
        bad = [x for x in items if not x[1]]
        failed += bool(bad)
        print(f"{'PASS' if not bad else 'FAIL'}  {case}  ({len(items) - len(bad)}/{len(items)} checks)")
        for what, ok, msg in bad:
            print(f"        x {what}: {msg}")
    print(f"\n{len(by_case) - failed}/{len(by_case)} cases passed, "
          f"{sum(1 for r in results if r[2])}/{len(results)} checks.")
    sys.exit(1 if failed else 0)


if __name__ == "__main__":
    main()
